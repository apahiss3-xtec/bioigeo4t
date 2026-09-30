"""Auditoria prova escrita <-> autoavaluació (4t ESO, SA2-SA7).

L'autoavaluació pre-examen (web/src/data/saN/avaluacio.js: blocs `escrita`,
`test` i `c`) ha de PREPARAR la prova sense copiar-la. Proves de referència:
SA2 té prova pròpia («Ostres de tres jocs», S4); SA3-SA7 es preparen amb el
bloc corresponent de la prova final del curs («El llac sota el gel», S29).
SA1 no té prova (`teProva: false`) i no s'audita.

Comprovacions:
  1. Cap seqüència de >= 8 paraules seguides coincident entre qualsevol text
     de la prova (B, C/A, solucionari) i qualsevol text de l'autoavaluació.
  2. Cap cas ni personatge coincident: marques del cas de la prova, noms
     propis de la prova (paraules amb majúscula enmig de frase) i xifres
     "de cas" (decimals o milers) repetides (aquestes, com a avís).
  3. L'escrita no reaprofita el cas del test ni l'enigma de la SA: cap
     seqüència de 8 paraules amb el test i cap marca de l'enigma o de les
     sessions (ENIGMA).
  4. Test de transferència i preguntes de la versió C: la correcta no sempre
     a la mateixa posició, ni la primera a totes, ni la més llarga en més d'una.
     La posició que es comprova és la que VEU l'alumne (TransferTest i
     AutoavaluacioC barregen les opcions amb permutacioEstable de src/utils.js).
  5. Cada `source` de l'escrita diu «Entrena el bloc ...».

Ús:  python scripts-avaluacio/audita_autoavaluacio.py [sa2 sa3 ...]   (des de web/)
Surt amb codi 1 si hi ha cap error (🔴); els avisos (🟡) no fan fallar.
"""
import html
import json
import re
import subprocess
import sys
import unicodedata
from pathlib import Path

WEB = Path(__file__).resolve().parents[1]
ROOT = WEB.parent

FINAL = ['SA7-origen-de-tot/S29-debat-i-prova-final/prova_final_curs_A.html',
         'SA7-origen-de-tot/S29-debat-i-prova-final/prova_final_curs_B.html',
         'SA7-origen-de-tot/S29-debat-i-prova-final/solucionari_prova_final_curs.html']

# Fitxers de prova per SA (glob relatius a l'arrel del projecte).
PROVES = {
    'sa2': ['SA2-cellula/S4-tancament-repas/prova_escrita_sa2_*.docx',
            'SA2-cellula/S4-tancament-repas/solucionari_prova_sa2.docx'],
    **{f'sa{n}': FINAL for n in range(3, 8)},
}

NGRAM = 8

# Elements que identifiquen el CAS de la prova. Cap d'ells pot aparèixer a
# l'autoavaluació (ni a l'escrita, ni al test, ni a la versió C).
LLAC = ['subglacial', 'sota el gel', 'antàrtida', 'perforació', 'perforar',
        'testimoni', 'testímoni', 'anticongelant', 'anf', 'europa', 'encèlad',
        'espores', '1,20', '0,80', '41 %', '87 %', '3.200', 'júpiter']
CAS = {
    'sa2': ['ostra', 'ostres', 'musclo', 'musclos', 'fangar', 'alfacs', 'laia',
            'youssef', 'triploide', 'triploides', 'tetraploide', 'lletosa',
            'lletoses', 'hatchery', 'crassostrea', 'bateas', 'batea',
            'neoplàsia', 'hemolimfa', '200 000', '200.000'],
    **{f'sa{n}': LLAC for n in range(3, 8)},
}

# Casos que l'ESCRITA tampoc no pot fer servir: el del test de la mateixa SA,
# l'enigma i els casos de les sessions (sN.js).
ENIGMA = {
    'sa2': ['ceba', 'plastilina', 'ratolí', 'ratolins', 'medul·la', 'òvul de 24',
            'cromosoma 21', 'donació'],
    'sa3': ['fenilcetonúria', 'pku', 'taló', 'pah', 'falciforme', 'crispr',
            'bessones', 'extracció'],
    'sa4': ['duchenne', 'daltonisme', 'hemofília', 'abo', 'grup sanguini',
            'moneda', 'monedes', 'heredity'],
    'sa5': ['gorongosa', 'elefant', 'elefants', 'ivori', 'ullals', 'arna',
            'arnes', 'bedoll', 'gallina', 'gallines', 'plataner', 'plataners',
            'antibiòtic', 'pesticida', 'pollastre'],
    'sa6': ['cardona', 'gondwana', 'ebre', 'muntanya de sal', 'wegener'],
    'sa7': ['alh84001', 'mart', 'meteorit', 'james webb', 'exoplaneta',
            'miller', 'urey', 'daus'],
}

# Paraules amb majúscula que no identifiquen cap cas (sigles, termes, rètols).
NO_CAS = set('''
OA AS AN AE NA ADN ARN ARNm OA-P1 OA-P2 OA-P3 OA-P4 IE Temple ESO SA Bloc
Versió Figura Fig Rúbrica Annex Resposta Encara Sé Explico Interpreto Relaciono
Connecto Justifico Argumento Detecto Identifico Reconec Avaluo Descric Llegeixo
Comparo Calculo Ordeno Proposo Dedueixo Resolc Ho Si Quin Quina Quins Quines Per
Com Què On Amb Sense Recorda Paraules Marca Ordena Explica Calcula Completa
Relaciona Decideix Imagina Proposa Escriu Indica Observa Llegeix Anomena
Justifica Dona Bona NOM COGNOMS CURS GRUP DATA NIVELL GLOBAL Biologia Geologia
BIOLOGIA GEOLOGIA Catalunya Sol Terra Ma Lamarck Darwin Punnett Mendel Nom Grup
Data Control Negatiu Positiu Mitosi Meiosi Objectiu Generació Pista Clica
Imprimir PDF Ctrl Espai Desfer Llicència CC BY-NC-SA Albert Pahissa Jordi
Domènech Adaptat Ordre Reconeixement NoComercial CompartirIgual
'''.split())


def norm(s):
    s = unicodedata.normalize('NFKD', s.lower())
    s = ''.join(c for c in s if not unicodedata.combining(c))
    return re.findall(r"[a-z0-9·']+", s)


def llegeix_docx(p):
    import docx
    d = docx.Document(p)
    out = [x.text for x in d.paragraphs]
    for t in d.tables:
        for r in t.rows:
            seen = set()
            for c in r.cells:
                if id(c._tc) not in seen:
                    seen.add(id(c._tc))
                    out.append(c.text)
    return out


def llegeix_html(p):
    s = p.read_text(encoding='utf-8')
    s = re.sub(r'(?is)<(script|style)\b.*?</\1>', ' ', s)
    s = re.sub(r'(?i)<br\s*/?>', ' ', s)
    s = re.sub(r'(?i)</?(p|div|li|tr|td|th|h\d|section|table|ul|ol|header|footer)\b[^>]*>', '\n', s)
    s = re.sub(r'<[^>]+>', ' ', s)
    s = html.unescape(s)
    return [re.sub(r'\s+', ' ', l).strip() for l in s.split('\n')]


def textos_prova(sa):
    fitxers, textos = [], []
    for pat in PROVES.get(sa, []):
        for p in sorted(ROOT.glob(pat)):
            if p in fitxers or '_prova' in p.parts or '_arxiu' in p.parts:
                continue
            fitxers.append(p)
            if p.suffix == '.docx':
                textos += llegeix_docx(p)
            elif p.suffix == '.html':
                textos += llegeix_html(p)
            else:
                textos += p.read_text(encoding='utf-8').splitlines()
    return fitxers, [t for t in textos if t.strip()]


def avaluacio(sa):
    f = (WEB / 'src' / 'data' / sa / 'avaluacio.js').as_uri()
    u = (WEB / 'src' / 'utils.js').as_uri()
    # Els components barregen les opcions amb permutacioEstable(id|text):
    # afegim `_shown`, la posició que VEU l'alumne.
    js = ("Promise.all([import(%r),import(%r)]).then(([m,u])=>{"
          "const o=Object.values(m)[0];"
          "const pos=q=>u.permutacioEstable(q.id+'|'+(q.text??''),q.options.length).indexOf(q.correct);"
          "for(const q of (o.test?.questions??[]))q._shown=pos(q);"
          "for(const q of (o.c?.preguntes??[]))q._shown=pos(q);"
          "process.stdout.write(JSON.stringify(o))})" % (f, u))
    r = subprocess.run(['node', '-e', js], capture_output=True, text=True,
                       encoding='utf-8', cwd=WEB)
    if r.returncode:
        sys.exit(f'No puc llegir {sa}/avaluacio.js:\n{r.stderr}')
    return json.loads(r.stdout)


def textos_escrita(a):
    out = []
    e = a.get('escrita') or {}
    out.append(('escrita.intro', e.get('intro', '')))
    for q in e.get('questions', []):
        out.append((f"escrita.{q['id']}.text", q['text']))
        out.append((f"escrita.{q['id']}.as", q['model']['as']))
        out.append((f"escrita.{q['id']}.ae", q['model']['ae']))
        out.append((f"escrita.{q['id']}.aeWhy", q.get('aeWhy', '')))
        out.append((f"escrita.{q['id']}.source", q.get('source', '')))
        out += [(f"escrita.{q['id']}.must", m) for m in q.get('must', [])]
    return [(l, x) for l, x in out if x]


def textos_test(a):
    out = []
    t = a.get('test') or {}
    out.append(('test.context', t.get('context', '')))
    for q in t.get('questions', []):
        out.append((f"test.{q['id']}.text", q['text']))
        out += [(f"test.{q['id']}.opt", o) for o in q['options']]
        out += [(f"test.{q['id']}.fb", v) for v in q.get('feedback', {}).values()]
    c = a.get('c') or {}
    for q in c.get('preguntes', []):
        out += [(f"c.{q['id']}", q.get(k, '')) for k in ('llegir', 'text')]
        out += [(f"c.{q['id']}.opt", o) for o in q['options']]
    k = c.get('completar')
    if k:
        out.append((f"c.{k['id']}", k.get('frase', '') + ' ' + k.get('llegir', '')))
    return [(l, x) for l, x in out if x]


def ngrams(words, n=NGRAM):
    return {' '.join(words[i:i + n]) for i in range(len(words) - n + 1)}


def noms_propis(linies):
    noms = set()
    for l in linies:
        for frase in re.split(r'[.!?:;«»"()\n]|—|·', l):
            paraules = frase.strip().split()
            for w in paraules[1:]:
                w = w.strip(",'’")
                if re.fullmatch(r"[A-ZÀÈÉÍÒÓÚÇ][a-zàèéíòóúïüç·]{2,}", w) and w not in NO_CAS:
                    noms.add(w)
    return noms


def xifres_cas(linies):
    s = ' '.join(linies)
    return set(re.findall(r'\b\d+,\d+\b|\b\d{1,3}\.\d{3}\b', s))


def busca_marques(marques, textos):
    baix = [(l, x.lower()) for l, x in textos]
    for marca in marques:
        rx = r'(?<!\w)' + re.escape(marca) + r'(?!\w)'
        on = sorted({l for l, x in baix if re.search(rx, x)})
        if on:
            yield marca, on


def comprova_opcions(etiqueta, qs, errors, avisos):
    if not qs:
        return
    pos = [q.get('_shown', q['correct']) for q in qs]
    llarga = []
    for q in qs:
        lens = [len(o) for o in q['options']]
        c = lens[q['correct']]
        if c == max(lens) and lens.count(c) == 1:
            llarga.append(q['id'])
    print(f"  {etiqueta}: posició a pantalla {[p + 1 for p in pos]} · la més llarga és la correcta a {llarga or 'cap'}")
    if len(qs) > 1 and len(set(pos)) == 1:
        errors.append(f"{etiqueta}: a pantalla la correcta és SEMPRE la posició {pos[0] + 1}")
    if len(qs) > 1 and pos.count(0) == len(qs):
        errors.append(f"{etiqueta}: a pantalla la correcta és sempre la primera")
    if len(llarga) > 1:
        errors.append(f"{etiqueta}: la correcta és la MÉS LLARGA a {llarga}")
    elif llarga:
        avisos.append(f"{etiqueta}: la correcta és la més llarga a {llarga} (només una, tolerat)")


def audita(sa):
    errors, avisos = [], []
    fitxers, prova = textos_prova(sa)
    a = avaluacio(sa)
    esc, tst = textos_escrita(a), textos_test(a)
    auto = esc + tst
    print(f"\n=== {sa.upper()} ===")
    if not fitxers:
        errors.append("no s'ha trobat la prova de referència")
    else:
        print('  prova:', ', '.join(f.name for f in fitxers))

    # 1 · frases coincidents amb la prova
    ng_prova = set()
    for l in prova:
        ng_prova |= ngrams(norm(l))
    for etiq, txt in auto:
        comuns = ngrams(norm(txt)) & ng_prova
        if comuns:
            errors.append(f"{etiq}: {len(comuns)} seqüència/es de {NGRAM} paraules de la prova, p. ex. «{sorted(comuns)[0]}»")

    # 2 · cas i personatges de la prova
    tot_auto = ' '.join(x for _, x in auto)
    tot_min = ' '.join(prova) + ' ' + tot_auto
    for nom in sorted(noms_propis(prova)):
        if re.search(r'(?<!\w)' + re.escape(nom.lower()) + r'(?!\w)', tot_min):
            continue
        if re.search(r'\b' + re.escape(nom) + r'\b', tot_auto):
            on = [l for l, x in auto if re.search(r'\b' + re.escape(nom) + r'\b', x)]
            errors.append(f"personatge o lloc de la prova «{nom}» a {', '.join(sorted(set(on))[:4])}")
    for marca, on in busca_marques(CAS.get(sa, []), auto):
        errors.append(f"element del cas de la prova «{marca}» a {', '.join(on[:4])}")
    for x in sorted(xifres_cas(prova)):
        if re.search(r'(?<![\d,.])' + re.escape(x) + r'(?![\d,])', tot_auto):
            avisos.append(f"la xifra «{x}» de la prova reapareix a l'autoavaluació (comprova que no sigui el mateix cas)")

    # 3 · l'escrita no reaprofita el test ni l'enigma
    ng_test = set()
    for _, x in tst:
        ng_test |= ngrams(norm(x))
    for etiq, txt in esc:
        comuns = ngrams(norm(txt)) & ng_test
        if comuns:
            errors.append(f"{etiq}: {len(comuns)} seqüència/es de {NGRAM} paraules del test, p. ex. «{sorted(comuns)[0]}»")
    for marca, on in busca_marques(ENIGMA.get(sa, []), esc):
        errors.append(f"l'escrita fa servir el cas del test o de les sessions «{marca}» a {', '.join(on[:4])}")

    # 4 · posicions i longituds
    comprova_opcions('test', (a.get('test') or {}).get('questions', []), errors, avisos)
    comprova_opcions('c.preguntes', (a.get('c') or {}).get('preguntes', []), errors, avisos)

    # 5 · sources
    for q in (a.get('escrita') or {}).get('questions', []):
        if not q.get('source', '').startswith('Entrena el bloc'):
            errors.append(f"escrita.{q['id']}.source no diu «Entrena el bloc …»: «{q.get('source')}»")

    for e in errors:
        print('  🔴', e)
    for w in avisos:
        print('  🟡', w)
    if not errors:
        print('  ✅ cap coincidència ni problema de posicions')
    return len(errors)


if __name__ == '__main__':
    sys.stdout.reconfigure(encoding='utf-8')
    sas = sys.argv[1:] or [f'sa{n}' for n in range(2, 8)]
    n = sum(audita(s) for s in sas)
    print(f"\n{'🔴 ' + str(n) + ' error/s' if n else '✅ Tot net'}")
    sys.exit(1 if n else 0)
